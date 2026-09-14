import base64
import io
import json
import mimetypes
import os
import re
import secrets
import time
import urllib.request
import urllib.error
import zipfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as xml_escape

from dotenv import load_dotenv
load_dotenv(Path(__file__).with_name('.env'))

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from google import genai
from google.genai import types

from gemini import chat, MODEL
from security import require_api_key, MAX_TEXT_CHARS

app = FastAPI(title='ZOMA AI API', version='1.3.0')
origins = [x.strip() for x in os.getenv('ALLOWED_ORIGINS', 'https://mohamkhamis209-hub.github.io,http://127.0.0.1:5500,http://localhost:5500').split(',') if x.strip()]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=False, allow_methods=['GET','POST','OPTIONS'], allow_headers=['*'])

MAX_UPLOAD = 20 * 1024 * 1024
IMAGE_MODEL = os.getenv('ZOMA_IMAGE_MODEL', 'gemini-3.1-flash-image')
TRANSFER_TTL = 600
transfer_store: dict[str, dict[str, Any]] = {}

class Message(BaseModel):
    role: str = Field(pattern='^(user|assistant)$')
    text: str = Field(min_length=1, max_length=MAX_TEXT_CHARS)
class ChatRequest(BaseModel):
    messages: list[Message] = Field(min_length=1, max_length=40)
class TransferCreate(BaseModel):
    payload: dict[str, Any]
class TransferReceive(BaseModel):
    code: str = Field(min_length=6, max_length=6)
class FileCreateRequest(BaseModel):
    file_type: str = Field(pattern='^(pdf|docx|xlsx|zip)$')
    prompt: str = Field(min_length=1, max_length=30000)
    filename: str | None = Field(default=None, max_length=120)


def cleanup_transfers():
    now = time.time()
    for code, item in list(transfer_store.items()):
        if item['expires'] <= now:
            transfer_store.pop(code, None)


def proxy_opener():
    proxy = os.getenv('https_proxy') or os.getenv('HTTPS_PROXY')
    if proxy:
        return urllib.request.build_opener(urllib.request.ProxyHandler({'http': proxy, 'https': proxy}))
    return urllib.request.build_opener()


def image_interaction(prompt: str, image_bytes: bytes | None = None, mime_type: str | None = None):
    """Use Gemini's current native image generation/editing Interactions API."""
    api_key = require_api_key()
    inputs: Any = prompt
    if image_bytes is not None:
        inputs = [
            {'type': 'text', 'text': prompt},
            {'type': 'image', 'mime_type': mime_type or 'image/png', 'data': base64.b64encode(image_bytes).decode('ascii')},
        ]
    body = {
        'model': IMAGE_MODEL,
        'input': inputs,
        'response_format': {'type': 'image', 'image_size': '1K'},
    }
    req = urllib.request.Request(
        'https://generativelanguage.googleapis.com/v1beta/interactions',
        data=json.dumps(body).encode('utf-8'),
        headers={'Content-Type': 'application/json', 'x-goog-api-key': api_key},
        method='POST',
    )
    try:
        with proxy_opener().open(req, timeout=180) as response:
            payload = json.loads(response.read().decode('utf-8'))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode('utf-8', 'ignore')[:1500]
        raise HTTPException(status_code=502, detail=f'خدمة الصور رفضت الطلب: {detail}') from exc
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f'تعذر الاتصال بخدمة الصور: {type(exc).__name__}') from exc

    # Support both the convenience output_image shape and the raw output blocks.
    image_data = None
    mime = 'image/png'
    text = ''
    oi = payload.get('output_image')
    if isinstance(oi, dict):
        image_data = oi.get('data')
        mime = oi.get('mime_type') or oi.get('mimeType') or mime
    for item in payload.get('output', []) or []:
        if not isinstance(item, dict):
            continue
        if item.get('type') in ('image', 'output_image'):
            image_data = image_data or item.get('data')
            mime = item.get('mime_type') or item.get('mimeType') or mime
        if item.get('type') == 'text':
            text += str(item.get('text') or '')
    if not image_data:
        # Some SDK/API versions expose a nested response object.
        for key in ('image', 'generated_image'):
            value = payload.get(key)
            if isinstance(value, dict) and value.get('data'):
                image_data = value['data']
                mime = value.get('mime_type') or value.get('mimeType') or mime
                break
    if not image_data:
        raise HTTPException(status_code=502, detail='لم ترجع خدمة الصور صورة. جرّب وصفًا آخر أو تحقق من تفعيل نموذج الصور في مفتاح Gemini.')
    return {'data': image_data, 'mime_type': mime, 'text': text.strip()}


@app.get('/')
def root():
    return {'name': 'ZOMA AI API', 'status': 'ok', 'model': MODEL, 'imageModel': IMAGE_MODEL}

@app.get('/api/health')
def health():
    return {'ok': True, 'gemini_configured': bool(os.getenv('GEMINI_API_KEY', '').strip()), 'model': MODEL, 'imageModel': IMAGE_MODEL}

@app.post('/api/chat')
def api_chat(payload: ChatRequest):
    try:
        return {'ok': True, 'reply': chat([m.model_dump() for m in payload.messages]), 'model': MODEL}
    except HTTPException:
        raise
    except Exception as exc:
        print('ZOMA CHAT ERROR:', type(exc).__name__, repr(exc), flush=True)
        raise HTTPException(status_code=500, detail=f'ZOMA ERROR: {type(exc).__name__}: {str(exc)}') from exc

@app.post('/api/analyze-file')
async def analyze_file(file: UploadFile = File(...), prompt: str = Form('اقرأ الملف أو حلل الصورة ونفذ المطلوب. أجب بالعربية.')):
    data = await file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(status_code=413, detail='حجم الملف أكبر من 20MB.')
    name = file.filename or 'file'
    content_type = file.content_type or mimetypes.guess_type(name)[0] or 'application/octet-stream'
    try:
        if content_type.startswith('image/'):
            client = genai.Client(api_key=require_api_key(), http_options=types.HttpOptions(client_args={'proxy': os.getenv('https_proxy') or os.getenv('HTTPS_PROXY') or 'http://proxy.server:3128'}, timeout=120000))
            try:
                response = client.models.generate_content(
                    model=MODEL,
                    contents=[types.Part.from_bytes(data=data, mime_type=content_type), prompt[:MAX_TEXT_CHARS]],
                    config=types.GenerateContentConfig(system_instruction='أنت ZOMA AI. حلل الصور بشكل مفيد وآمن، وإذا كانت الصورة لسؤال دراسي فاشرح الحل خطوة بخطوة. أجب بالعربية.', temperature=0.5, max_output_tokens=4096),
                )
                return {'ok': True, 'reply': (response.text or '').strip()}
            finally:
                try: client.close()
                except Exception: pass
        if content_type == 'application/pdf' or name.lower().endswith('.pdf'):
            client = genai.Client(api_key=require_api_key(), http_options=types.HttpOptions(client_args={'proxy': os.getenv('https_proxy') or os.getenv('HTTPS_PROXY') or 'http://proxy.server:3128'}, timeout=120000))
            try:
                response = client.models.generate_content(
                    model=MODEL,
                    contents=[types.Part.from_bytes(data=data, mime_type='application/pdf'), prompt[:MAX_TEXT_CHARS]],
                    config=types.GenerateContentConfig(system_instruction='أنت ZOMA AI. اقرأ الملف المرفق واستخرج المعلومات المهمة ونفذ طلب المستخدم. أجب بالعربية.', temperature=0.4, max_output_tokens=4096),
                )
                return {'ok': True, 'reply': (response.text or '').strip()}
            finally:
                try: client.close()
                except Exception: pass
        if name.lower().endswith('.docx') or content_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                xml = z.read('word/document.xml').decode('utf-8', 'ignore')
            text = re.sub(r'<[^>]+>', ' ', xml)
            text = ' '.join(text.split())[:MAX_TEXT_CHARS]
        elif name.lower().endswith('.zip') or content_type in ('application/zip', 'application/x-zip-compressed'):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                infos = [i for i in z.infolist() if not i.is_dir()]
                parts = [f'محتويات ZIP ({len(infos)} ملف):']
                for i in infos[:200]:
                    parts.append(f'- {i.filename} ({i.file_size} bytes)')
                    lower = i.filename.lower()
                    if i.file_size <= 500_000 and lower.endswith(('.txt','.md','.json','.csv','.html','.css','.js','.py','.xml')):
                        try: parts.append(z.read(i).decode('utf-8', 'ignore')[:8000])
                        except Exception: pass
                text = '\n'.join(parts)[:MAX_TEXT_CHARS]
        else:
            is_text = content_type.startswith('text/') or name.lower().endswith(('.txt','.md','.json','.csv','.html','.css','.js','.py','.xml'))
            if not is_text:
                raise HTTPException(status_code=415, detail='هذا النوع من الملفات غير مدعوم حاليًا. جرّب PDF أو DOCX أو ZIP أو ملفًا نصيًا.')
            text = data.decode('utf-8', 'ignore')[:MAX_TEXT_CHARS]
        r = chat([{'role': 'user', 'text': (prompt or 'اقرأ الملف ونفذ المطلوب.') + '\n\nمحتوى الملف:\n' + text}])
        return {'ok': True, 'reply': r}
    except HTTPException:
        raise
    except zipfile.BadZipFile as exc:
        raise HTTPException(status_code=400, detail='ملف ZIP تالف أو غير صالح.') from exc
    except Exception as exc:
        print('ZOMA FILE ERROR:', type(exc).__name__, repr(exc), flush=True)
        raise HTTPException(status_code=502, detail=f'تعذر تحليل الملف الآن: {type(exc).__name__}') from exc

@app.post('/api/generate-image')
async def generate_image(prompt: str = Form(...), file: UploadFile | None = File(default=None)):
    prompt = prompt.strip()[:MAX_TEXT_CHARS]
    if not prompt:
        raise HTTPException(status_code=400, detail='اكتب وصف الصورة أو التعديل أولًا.')
    image_bytes = None
    mime = None
    if file is not None:
        image_bytes = await file.read(MAX_UPLOAD + 1)
        if len(image_bytes) > MAX_UPLOAD:
            raise HTTPException(status_code=413, detail='حجم الصورة أكبر من 20MB.')
        mime = file.content_type or mimetypes.guess_type(file.filename or '')[0] or 'image/png'
        if not mime.startswith('image/'):
            raise HTTPException(status_code=415, detail='الملف المرفوع يجب أن يكون صورة.')
        prompt = 'عدّل الصورة المرفقة حسب طلب المستخدم. حافظ على العناصر غير المطلوبة للتغيير، واجعل النتيجة طبيعية وعالية الجودة. طلب المستخدم: ' + prompt
    result = image_interaction(prompt, image_bytes, mime)
    return {'ok': True, 'data': result['data'], 'mime_type': result['mime_type'], 'text': result['text'], 'model': IMAGE_MODEL, 'edited': image_bytes is not None}


def generate_file_content(prompt: str) -> str:
    return chat([{'role': 'user', 'text': 'أنشئ محتوى جاهزًا ومنظمًا بناءً على الطلب التالي. لا تكتب مقدمات عن أنك نموذج ذكاء اصطناعي. استخدم العربية عند طلب المستخدم العربية. أعد المحتوى فقط.\n\n' + prompt[:MAX_TEXT_CHARS]}])


def safe_filename(name: str | None, fallback: str) -> str:
    raw = (name or fallback).strip()
    raw = re.sub(r'[^\w\-\u0600-\u06FF .]', '_', raw).strip() or fallback
    return raw[:80]


def make_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        def shape(s): return get_display(arabic_reshaper.reshape(s))
    except Exception:
        def shape(s): return s
    font_candidates = [
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/truetype/freefont/FreeSans.ttf',
    ]
    font_path = next((p for p in font_candidates if os.path.exists(p)), None)
    if not font_path:
        raise RuntimeError('لم يتم العثور على خط Unicode لإنشاء PDF.')
    pdfmetrics.registerFont(TTFont('ZomaUnicode', font_path))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45)
    styles = getSampleStyleSheet()
    style = ParagraphStyle('ZomaArabic', parent=styles['Normal'], fontName='ZomaUnicode', fontSize=12, leading=20, alignment=TA_RIGHT, spaceAfter=8)
    story = []
    for para in text.split('\n'):
        clean = para.strip()
        if clean:
            story.append(Paragraph(xml_escape(shape(clean)), style))
        else:
            story.append(Spacer(1, 7))
    doc.build(story)
    return buf.getvalue()


def make_docx(text: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    for line in text.split('\n'):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx(text: str) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'ZOMA AI'
    for row_no, line in enumerate(text.splitlines() or [''], 1):
        ws.cell(row=row_no, column=1, value=line)
    ws.column_dimensions['A'].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_zip(text: str, prompt: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('zoma-content.txt', text)
        z.writestr('README.txt', 'تم إنشاء هذا الملف بواسطة ZOMA AI.\n\nالطلب الأصلي:\n' + prompt[:5000])
    return buf.getvalue()

@app.post('/api/create-file')
def create_file(req: FileCreateRequest):
    try:
        content = generate_file_content(req.prompt)
        ext = req.file_type
        base = safe_filename(req.filename, 'zoma-document')
        if base.lower().endswith('.' + ext):
            filename = base
        else:
            filename = base + '.' + ext
        if ext == 'pdf':
            data = make_pdf(content); mime = 'application/pdf'
        elif ext == 'docx':
            data = make_docx(content); mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        elif ext == 'xlsx':
            data = make_xlsx(content); mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:
            data = make_zip(content, req.prompt); mime = 'application/zip'
        return {'ok': True, 'filename': filename, 'mime_type': mime, 'data': base64.b64encode(data).decode('ascii')}
    except HTTPException:
        raise
    except Exception as exc:
        print('ZOMA CREATE FILE ERROR:', type(exc).__name__, repr(exc), flush=True)
        raise HTTPException(status_code=502, detail=f'تعذر إنشاء الملف الآن: {type(exc).__name__}: {str(exc)}') from exc

@app.post('/api/transfer/create')
def transfer_create(req: TransferCreate):
    cleanup_transfers()
    if not req.payload: raise HTTPException(status_code=400, detail='بيانات النقل فارغة.')
    for _ in range(30):
        code = f'{secrets.randbelow(1000000):06d}'
        if code not in transfer_store: break
    else: raise HTTPException(status_code=503, detail='تعذر إنشاء كود النقل الآن.')
    transfer_store[code] = {'payload': req.payload, 'expires': time.time() + TRANSFER_TTL}
    return {'ok': True, 'code': code, 'expiresIn': TRANSFER_TTL}

@app.post('/api/transfer/receive')
def transfer_receive(req: TransferReceive):
    cleanup_transfers()
    item = transfer_store.pop(req.code.strip(), None)
    if not item: raise HTTPException(status_code=404, detail='كود النقل غير موجود أو انتهت صلاحيته.')
    return {'ok': True, 'payload': item['payload']}
